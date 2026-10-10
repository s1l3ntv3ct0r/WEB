from pathlib import Path
from copy import deepcopy
import json
import re
from datetime import datetime, timezone
from docx import Document
from docx.shared import Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT

ROOT = Path(__file__).resolve().parents[2]
WORK = Path(__file__).resolve().parent
SOURCE = ROOT / 'pentest/reports/ShopLab_BaoCaoTongHop_Phan1_Den8.docx'
OUTPUT = ROOT / 'pentest/reports/ShopLab_BaoCaoTongHop_HoanChinh.docx'

def set_text(p, text):
    runs = p.runs
    template = deepcopy(runs[0]._r.rPr) if runs and runs[0]._r.rPr is not None else None
    for child in list(p._p):
        if child.tag != qn('w:pPr'):
            p._p.remove(child)
    r = p.add_run(text)
    if template is not None:
        r._r.insert(0, template)

def normalize_run(run):
    rp = run.find(qn('w:rPr'))
    if rp is None:
        rp = OxmlElement('w:rPr')
        run.insert(0, rp)
    for tag, attrs in [
        ('rFonts', {'ascii': 'Times New Roman', 'hAnsi': 'Times New Roman', 'eastAsia': 'Times New Roman', 'cs': 'Times New Roman'}),
        ('sz', {'val': '26'}), ('szCs', {'val': '26'}), ('color', {'val': '000000'}),
    ]:
        for old in list(rp.findall(qn('w:' + tag))):
            rp.remove(old)
        el = OxmlElement('w:' + tag)
        for key, value in attrs.items():
            el.set(qn('w:' + key), value)
        rp.append(el)

def add_captcha(doc, anchor):
    evidence = ROOT / 'pentest/runs/captcha-assessment-20261006-r2'
    results = json.loads((evidence / 'results.json').read_text(encoding='utf-8'))
    assert results['success'] and results['sourceUnchanged']
    cases = []
    for suite in ['lifecycle', 'isolation', 'validation', 'rate']:
        data = json.loads((evidence / (suite + '.json')).read_text(encoding='utf-8'))
        assert data['success'] and all(c['passed'] for c in data['cases'])
        cases.extend(data['cases'])
    assert len(cases) == 15 and sum(s['records'] for s in results['suites']) == 58
    initial = set(doc.element.body)
    def p(text, style='Normal'):
        return doc.add_paragraph(text, style)
    title = p('BỔ SUNG KIỂM THỬ CAPTCHA ĐĂNG NHẬP', 'Heading 1')
    title.paragraph_format.page_break_before = True
    p('Thực hiện ngày 06/10/2026 trên mã nguồn ShopLab trong thư mục dự án. Phần bổ sung đánh giá cơ chế CAPTCHA được thêm sau đợt kiểm thử ngày 25–26/09/2026. Hai endpoint được kiểm tra là POST /api/auth/login và POST /api/auth/login-captcha.')
    p('Kết quả: 2 bài kiểm thử tích hợp và 15 tình huống bổ sung đạt tiêu chí. Bộ tình huống bổ sung ghi 58 request/response HTTP. CAPTCHA được kiểm tra tại backend; gửi đúng mật khẩu trực tiếp đến API khi thiếu mã bắt buộc vẫn bị từ chối. Không ghi nhận khả năng bỏ qua CAPTCHA trong các tình huống đã thực hiện.')
    p('Cơ chế xác minh và phạm vi kiểm thử', 'Heading 2')
    p('Email được loại khoảng trắng đầu cuối và chuyển về chữ thường. Bộ đếm lưu theo SHA-256 của email trong bảng LoginAttempt. Sau lần nhập sai thứ ba, API trả HTTP 401 LOGIN_FAILED kèm ảnh CAPTCHA. Từ lần tiếp theo, yêu cầu phải có captchaId và captchaAnswer hợp lệ trước khi backend kiểm tra mật khẩu.')
    p('Ảnh CAPTCHA gồm 6 ký tự; ID được tạo từ 24 byte ngẫu nhiên. Đáp án được chuẩn hóa bằng trim và chuyển chữ hoa trước khi đối chiếu hash ràng buộc với email và ID. Mã có thời hạn 2 phút. Trạng thái sai có thời hạn 15 phút, được gia hạn khi phát sinh thêm lần kiểm tra mật khẩu thất bại. Đăng nhập thành công xóa trạng thái của chính email đó.')
    p('Giới hạn theo IP là 30 request trong 15 phút và được dùng chung cho register, login, login-captcha, forgot-password và reset-password. Bộ giới hạn chạy trước kiểm tra dữ liệu và CAPTCHA; vì vậy response 429 là chặn theo hạn mức IP, còn 401 CAPTCHA_REQUIRED là thiếu hoặc không đáp ứng mã xác minh của một email.')
    p('Phương pháp và dữ liệu thử nghiệm', 'Heading 2')
    p('Các request được gửi bằng Supertest vào cùng Express app của dự án, sử dụng PostgreSQL và cơ sở dữ liệu thử nghiệm riêng. Mỗi nhóm chạy trong một process mới để tách bộ đếm rate limit. Tài khoản thử nghiệm được tạo riêng; không dùng dữ liệu tài khoản đang có trong cơ sở dữ liệu phát triển.')
    p('Để đối chiếu trường hợp mã đúng và hết hạn, một số phép thử đặt hash của mã biết trước hoặc timestamp hết hạn trong cơ sở dữ liệu thử nghiệm. Các phép thử này xác minh logic kiểm tra tại backend; không giải ảnh bằng OCR và không đo thời gian chờ thực 2 hoặc 15 phút. Giao diện được rà soát từ AuthPage.jsx và đối chiếu ảnh thực nghiệm ngày 01/10/2026.')
    p('Kết quả chi tiết', 'Heading 2')
    rows = [
        ('CAP-01', 'Nhập sai cùng email 3 lần', 'Lần 1–2: 401 LOGIN_FAILED, chưa có CAPTCHA. Lần 3: 401 kèm ảnh PNG; failures = 3.'),
        ('CAP-02', 'Mật khẩu đúng nhưng thiếu CAPTCHA bắt buộc', '401 CAPTCHA_REQUIRED; không có accessToken hoặc Set-Cookie.'),
        ('CAP-03', 'Gửi đáp án sai', '401 CAPTCHA_REQUIRED; ID ảnh đổi, không cấp cookie phiên.'),
        ('CAP-04', 'Mã hết hạn nhưng bộ đếm còn hiệu lực', '401 CAPTCHA_REQUIRED. Timestamp hết hạn được đặt trong DB thử nghiệm.'),
        ('CAP-05', 'CAPTCHA đúng, mật khẩu sai rồi dùng lại mã', 'Lần đầu: 401 LOGIN_FAILED và đổi ID. Gửi lại mã cũ với mật khẩu đúng: 401 CAPTCHA_REQUIRED.'),
        ('CAP-06', 'Xin mã mới rồi gửi mã trước đó', 'Đổi mã: 200 và ID mới. Mã cũ bị 401 CAPTCHA_REQUIRED.'),
        ('CAP-07', 'Đáp án chữ thường có khoảng trắng và mật khẩu đúng', '200, có token và cookie; trạng thái email bị xóa. Lần sai tiếp theo chưa có CAPTCHA.'),
        ('CAP-08', 'Đăng nhập B đúng khi A đang bắt buộc CAPTCHA', 'B nhận 200. A vẫn failures = 3 và bị 401 CAPTCHA_REQUIRED khi thiếu mã.'),
        ('CAP-09', 'Dùng mã của A cho email B', '401 CAPTCHA_REQUIRED; mã không dùng được giữa hai email.'),
        ('CAP-10', 'Đổi email thành chữ hoa và thêm khoảng trắng', 'Vẫn 401 CAPTCHA_REQUIRED, cùng trạng thái failures = 3.'),
        ('CAP-11', 'Ba lần sai gửi đồng thời', 'Cả ba trả 401; một response kèm ảnh, DB ghi đủ failures = 3.'),
        ('CAP-12', 'Trạng thái sai hết hạn, mật khẩu đúng', '200 không cần CAPTCHA; trạng thái email bị xóa. Timestamp được đặt trong DB thử nghiệm.'),
        ('CAP-13', 'Xin ảnh trước ngưỡng cho email tồn tại và chưa đăng ký', 'Cả hai trả 200 với captcha = null; riêng phép thử này chưa phân biệt tài khoản bằng trạng thái CAPTCHA.'),
        ('CAP-14', 'ID dài, đáp án dài, thêm role hoặc email sai', 'Bốn request trả 400 VALIDATION_ERROR; không cấp cookie, không tạo LoginAttempt.'),
        ('CAP-15', 'Dùng hết quota qua route xin ảnh rồi đăng nhập đúng', '30 request đầu: 200. Request 31: 429 RATE_LIMIT. Đăng nhập đúng tiếp theo vẫn 429; có RateLimit và Retry-After.'),
    ]
    t = doc.add_table(rows=1, cols=3)
    t.style = 'Table Grid'
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    from docx.shared import Cm
    widths = [Cm(2.1), Cm(6.0), Cm(8.9)]
    for column, width in zip(t.columns, widths):
        column.width = width
    for c, text in zip(t.rows[0].cells, ['Mã', 'Tình huống', 'Kết quả ghi nhận']):
        c.text = text
    for row in rows:
        for c, text in zip(t.add_row().cells, row):
            c.text = text
    borders = OxmlElement('w:tblBorders')
    for edge in ['top', 'left', 'bottom', 'right', 'insideH', 'insideV']:
        border = OxmlElement('w:' + edge)
        for key, value in {'val': 'single', 'sz': '4', 'color': 'D9D9D9'}.items():
            border.set(qn('w:' + key), value)
        borders.append(border)
    pr = t._tbl.tblPr
    for old in list(pr.findall(qn('w:tblBorders'))):
        pr.remove(old)
    pr.append(borders)
    for i, row in enumerate(t.rows):
        trpr = row._tr.get_or_add_trPr()
        trpr.append(OxmlElement('w:cantSplit'))
        if i == 0:
            trpr.append(OxmlElement('w:tblHeader'))
        for j, (cell, width) in enumerate(zip(row.cells, widths)):
            cell.width = width
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            tcpr = cell._tc.get_or_add_tcPr()
            margins = OxmlElement('w:tcMar')
            for edge, size in [('top', '70'), ('bottom', '70'), ('left', '85'), ('right', '85')]:
                el = OxmlElement('w:' + edge); el.set(qn('w:w'), size); el.set(qn('w:type'), 'dxa'); margins.append(el)
            tcpr.append(margins)
            if i == 0:
                shade = OxmlElement('w:shd'); shade.set(qn('w:fill'), 'E7E6E6'); tcpr.append(shade)
            for para in cell.paragraphs:
                para.paragraph_format.space_before = Pt(2)
                para.paragraph_format.space_after = Pt(2)
                para.paragraph_format.line_spacing = 1.0
                para.alignment = WD_ALIGN_PARAGRAPH.CENTER if j == 0 else WD_ALIGN_PARAGRAPH.LEFT
                for run in para.runs:
                    run.font.size = Pt(13); run.bold = i == 0
    p('Bằng chứng HTTP của CAP-01 đến CAP-07 nằm trong lifecycle.json; CAP-08 đến CAP-12 trong isolation.json; CAP-13 và CAP-14 trong validation.json; CAP-15 trong rate.json. Mỗi bản ghi có mã E-CAP, thời điểm, endpoint, status, mã lỗi, trạng thái CAPTCHA và dấu hiệu cấp token/cookie; mật khẩu và đáp án thử nghiệm đã được che.')
    p('Minh chứng CAPTCHA trên giao diện', 'Heading 2')
    para = doc.add_paragraph(); para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    para.paragraph_format.keep_with_next = True
    para.add_run().add_picture(str(ROOT / 'report_work/bruteforce_source/image9.png'), width=Cm(15.5))
    cap = p('Hình CAP 01 Form đăng nhập hiển thị ảnh CAPTCHA và ô nhập ký tự trong thực nghiệm ngày 01/10/2026', 'Caption'); cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p('Ảnh giao diện thể hiện yêu cầu hoàn thành CAPTCHA, ảnh 6 ký tự, nút Get a new code và ô nhập ký tự. Log ngày 01/10/2026 ghi lần sai thứ ba của cùng email có CAPTCHA; sau một lần đăng nhập đúng ở tài khoản test, email đó vẫn nhận CAPTCHA_REQUIRED. Kết quả bổ sung ngày 06/10/2026 xác nhận lại tính độc lập của bộ đếm theo email.')
    p('Đối chiếu thử đăng nhập bằng Burp Intruder', 'Heading 2')
    p('Danh sách trong thực nghiệm Burp có 300 cặp email và mật khẩu được chuẩn bị. Lượt bị giới hạn theo IP đi đến cặp thứ 28 thì nhận 429; cặp này bị chặn trước kiểm tra mật khẩu. Số 300 là quy mô dữ liệu đầu vào, không phải số cặp đã hoàn tất xác thực. Một ảnh của lượt khác ghi nhận cặp admin hợp lệ trả 200; hai kết quả không thuộc cùng chuỗi request.')
    p('Việc đăng nhập thành công bằng cặp thông tin hợp lệ không chứng minh vượt CAPTCHA. Để kết luận bỏ qua cơ chế này, cùng email phải đang ở trạng thái bắt buộc CAPTCHA nhưng vẫn được cấp phiên khi thiếu hoặc sai mã. CAP-02, CAP-03, CAP-05 và CAP-08 đều ghi nhận backend từ chối các trường hợp tương ứng. CAPTCHA đúng cũng không thay thế việc kiểm tra mật khẩu.')
    p('Quan sát về việc đổi mã giữa các lượt đăng nhập', 'Heading 2')
    p('O-CAP-001 — Informational: challenge được lưu chung theo email và endpoint xin mã mới không yêu cầu phiên đăng nhập. CAP-06 xác nhận request đổi mã trả 200 và làm mã trước đó mất hiệu lực. Một lượt xin mã khác cho cùng email vì vậy có thể làm người đang nhập phải lấy mã mới. Phép thử chỉ ghi nhận một lần thay mã; chưa đo khả năng gây gián đoạn liên tục hoặc tấn công từ nhiều nguồn IP.')
    p('Khi mở rộng triển khai, cần cân nhắc ràng buộc challenge với từng lượt xác thực và giới hạn riêng hành vi xin mã mới. Kiểm tra lại phải bảo đảm mã của một lượt không bị thay bởi lượt khác ngoài thiết kế, đồng thời không cho phép tái sử dụng mã đã hết hạn hoặc đã dùng. Quan sát này được theo dõi riêng với bốn lỗ hổng Low của đợt gốc.')
    p('Kết luận và bảo toàn dữ liệu', 'Heading 2')
    p('Các trường hợp đã kiểm tra cho thấy CAPTCHA được thực thi tại API, mã cũ hoặc hết hạn bị từ chối, đăng nhập thành công đặt lại đúng bộ đếm và ba lần sai đồng thời được ghi đầy đủ. Giới hạn IP vẫn có hiệu lực với route xin mã và đăng nhập đúng. Kết quả chưa đánh giá OCR, password spraying phân tán, nhiều replica, proxy production hoặc khả năng tiếp cận của CAPTCHA.')
    p('Cơ sở dữ liệu thử nghiệm đã được xóa và cluster thử nghiệm đã dừng sau lượt kiểm tra. Hash SHA-256 của sáu tệp nguồn và kiểm thử được theo dõi khớp trước/sau; mã nguồn ứng dụng không bị thay đổi. Hồ sơ bổ sung lưu tại pentest/runs/captcha-assessment-20261006-r2, gồm results.json, jest-results.json, bốn log HTTP và captcha-api.png. Log ngày 01/10 lưu tại pentest/runs/captcha-demo-20261001/results.json; ảnh Burp và giao diện nằm trong báo cáo demo Brute force CAPTCHA.')
    p('Nguồn mã đối chiếu: backend/src/modules/auth/auth.routes.js; backend/src/modules/auth/login-captcha.js; frontend/src/pages/auth/AuthPage.jsx; backend/prisma/schema.prisma; migration 202609300001_login_captcha; backend/tests/integration/login-captcha.test.js.')
    added = [el for el in doc.element.body if el not in initial and el.tag != qn('w:sectPr')]
    for el in added:
        anchor.addprevious(el)

def main():
    doc = Document(SOURCE)
    original = list(doc.paragraphs)
    patch = json.loads((WORK / 'editorial_patch.json').read_text(encoding='utf-8'))
    replacements = {int(i): t for i, t in patch['replace'].items()}
    replacements.update({
        1: 'TỔNG HỢP PHẦN 1 ĐẾN PHẦN 8 VÀ KIỂM THỬ BỔ SUNG CAPTCHA',
        2: 'Đợt kiểm thử SL-PT-20260925 • Phiên bản báo cáo 1.1 • Cập nhật ngày 06/10/2026',
        4: 'Báo cáo trình bày kết quả kiểm thử ứng dụng ShopLab trong môi trường local, gồm khảo sát bề mặt, xác thực và phiên, phân quyền, logic nghiệp vụ, xử lý đầu vào, tệp và cấu hình triển khai. Đợt gốc diễn ra ngày 25–26/09/2026; hồ sơ được đối chiếu ngày 30/09/2026. Kiểm thử bổ sung CAPTCHA đăng nhập được thực hiện ngày 06/10/2026 trên mã nguồn hiện tại.',
        6: 'Đợt gốc xác nhận bốn phát hiện mức Low: phân biệt email khi đăng ký, checkout vượt giới hạn tiền, ID ngoài Int32 gây lỗi máy chủ và phạm vi phục vụ tệp của Vite. Cả bốn còn Open, chưa được kiểm tra lại sau sửa. Phần bổ sung CAPTCHA có 2 bài kiểm thử tích hợp và 15 tình huống đạt tiêu chí; không ghi nhận bỏ qua mã xác minh trong các trường hợp đã thử. Một quan sát Informational về việc đổi mã theo email được bổ sung. Docker/Nginx/TLS và image runtime chưa được kiểm thử thực tế.',
        20: 'Mỗi phép thử có mã tình huống và mã bằng chứng để tra cứu trong hồ sơ kỹ thuật. Đường dẫn tương đối trong báo cáo được tính từ thư mục dự án tại thời điểm kiểm thử.',
        24: 'PASS là đáp ứng tiêu chí của phép thử; REVIEW là kết quả cần đối chiếu hoặc xác minh thêm; PARTIAL là phép thử chưa thực hiện đầy đủ. Open là phát hiện chưa được khắc phục và kiểm tra lại. Số phát hiện được thống kê theo nguyên nhân, không theo số request hay số ảnh.',
        26: 'Khối lượng kiểm thử được thống kê riêng theo phương pháp thu thập của từng chương. Các bộ log HTTP client và traffic nền của trình duyệt không có cùng cách đếm nên không cộng thành tổng số request. Hồ sơ gốc gồm 30 ảnh bằng chứng.',
        939: 'Hồ sơ tổng kết gồm integrity-check.json đối chiếu manifest, merge-provenance.json ghi nguồn các báo cáo, findings-register.json thống kê phát hiện và RETEST_PLAN.md nêu tiêu chí kiểm tra lại sau sửa.',
        908: 'Bốn phát hiện duy nhất của đợt gốc là F-P2-001, F-P4-001, F-P5-001 và F-P7-001. Các trường hợp tái hiện cùng nguyên nhân được giữ chung một mã phát hiện. Năm quan sát Informational được theo dõi riêng: O-P6-001, O-P7-001/002/003 và O-CAP-001 về việc mã CAPTCHA được thay chung theo email. Phần bổ sung ngày 06/10/2026 không phát sinh lỗ hổng bỏ qua CAPTCHA trong các phép thử đã thực hiện.',
        936: 'Kết quả audit dependency thuộc ngày 26/09/2026. Kiểm thử bổ sung CAPTCHA ngày 06/10/2026 không chạy lại audit hoặc kiểm tra lại bốn phát hiện cũ. Các so sánh dữ liệu và source chỉ áp dụng cho từng lượt được ghi trong hồ sơ.',
        938: 'Hồ sơ kỹ thuật đợt gốc gồm CASE_LOG, HTTP JSON, đối chiếu DB/tệp, mã chạy kiểm thử, ảnh gốc và evidence manifest trong pentest/runs/SL-PT-20260925/part1 đến part7. Hồ sơ CAPTCHA bổ sung nằm tại pentest/runs/captcha-assessment-20261006-r2; log đối chứng trước đó tại pentest/runs/captcha-demo-20261001. Thông tin xác thực nguyên văn và backup DB được lưu riêng, không đính kèm báo cáo.',
        932: 'Production profile phải tắt seed, từ chối placeholder và dùng HTTPS/Secure đúng thiết kế. Kiểm tra hai client sau Nginx, header giả từ nguồn không tin cậy và audit dependency/image tại ngày kiểm tra lại. Tiêu chí PDF cần bộ mẫu hợp lệ và lỗi cấu trúc; magic bytes không xác nhận tệp sạch malware. Với O-CAP-001, kiểm tra việc tách challenge theo lượt xác thực, giữ bộ đếm sai theo email và từ chối mã hết hạn hoặc đã dùng. Cần đối chiếu hai trình duyệt cùng email để xác minh hành vi đổi mã theo thiết kế.',
    })
    deleted = set(patch['delete'])
    # Remove duplicated report titles from chapters; their Heading 1 already names the subject.
    deleted.update([29, 30, 69, 233, 381, 510, 649, 782, 940])
    changes = []
    for i, p in enumerate(original):
        if i in deleted:
            changes.append({'paragraph': i, 'action': 'delete', 'before': p.text})
            p._p.getparent().remove(p._p)
            continue
        text = replacements.get(i, p.text)
        # Redundant verdicts repeat the concrete observations directly below them.
        if p.style.name == 'Normal' and text.strip() == 'Kết luận: Đạt các tiêu chí đã thử.':
            changes.append({'paragraph': i, 'action': 'delete', 'before': p.text})
            p._p.getparent().remove(p._p)
            continue
        for phrase in [' Đạt tiêu chí đã thử.', ' PASS theo tiêu chí đã thử.']:
            text = text.replace(phrase, '')
        text = text.replace('Quan sát và giới hạn:', 'Kết quả:')
        if p.style.name.startswith('Heading'):
            text = re.sub(r'\s*[—–]\s*', ' ', text)
            text = re.sub(r'\b(F|P|O)-(P?\d+)-(\d+)\b', r'\1 \2 \3', text)
        if text != p.text:
            changes.append({'paragraph': i, 'action': 'replace', 'before': p.text, 'after': text})
            set_text(p, text)
    for change in patch['table_replace']:
        cell = doc.tables[change['table']].cell(change['row'], change['col'])
        changes.append({'table': change['table'], 'row': change['row'], 'col': change['col'], 'before': cell.text, 'after': change['text']})
        cell.text = change['text']
    # The supplement sits after the historical test chapters and before the final conclusions.
    add_captcha(doc, original[905]._p)
    # Update the scope summary without merging unrelated request counts.
    for table_index in [1, 22]:
        row = doc.tables[table_index].add_row()
        values = ['CAPTCHA', '06/10/2026: 2 bài kiểm thử tích hợp; 15 tình huống; 58 HTTP record', 'Instance và DB kiểm thử riêng; O-CAP-001 Informational'] if table_index == 1 else ['CAPTCHA đăng nhập', 'Ngưỡng sai, mã thiếu/sai/hết hạn, dùng lại, song song và rate limit', 'O-CAP-001; chưa đánh giá OCR hoặc tấn công phân tán']
        for cell, text in zip(row.cells, values):
            cell.text = text
    row = doc.tables[-1].add_row()
    for cell, text in zip(row.cells, ['Theo thiết kế đăng nhập', 'O-CAP-001 tách challenge theo lượt xác thực và giữ giới hạn xin mã', 'Người phụ trách backend và xác thực']):
        cell.text = text
    from docx.shared import Cm
    scope_widths = [Cm(2.8), Cm(6.5), Cm(7.7)]
    for column, width in zip(doc.tables[1].columns, scope_widths):
        column.width = width
    for row in doc.tables[1].rows:
        for cell, width in zip(row.cells, scope_widths):
            cell.width = width
    # Complete borders on appended rows and keep table headings consistent.
    for table in doc.tables:
        borders = OxmlElement('w:tblBorders')
        for edge in ['top', 'left', 'bottom', 'right', 'insideH', 'insideV']:
            el = OxmlElement('w:' + edge)
            for key, value in {'val': 'single', 'sz': '4', 'color': 'D9D9D9'}.items():
                el.set(qn('w:' + key), value)
            borders.append(el)
        for old in list(table._tbl.tblPr.findall(qn('w:tblBorders'))):
            table._tbl.tblPr.remove(old)
        table._tbl.tblPr.append(borders)
        for i, row in enumerate(table.rows):
            for cell in row.cells:
                if cell.text.strip() == 'Anonymous':
                    cell.text = 'Ẩn danh'
                    for paragraph in cell.paragraphs:
                        for run in paragraph.runs:
                            run.bold = True
                elif cell.text.strip() == 'Attachment/message':
                    cell.text = 'Tệp mỗi tin nhắn'
                tcpr = cell._tc.get_or_add_tcPr()
                for border in list(tcpr.findall(qn('w:tcBorders'))):
                    tcpr.remove(border)
                if i == 0:
                    for shade in list(tcpr.findall(qn('w:shd'))):
                        tcpr.remove(shade)
                    shade = OxmlElement('w:shd'); shade.set(qn('w:fill'), 'E7E6E6'); tcpr.append(shade)
    # Keep captions on the page containing their evidence image.
    for p in doc.paragraphs:
        if p._p.xpath('.//w:drawing'):
            p.paragraph_format.keep_with_next = True
            p.paragraph_format.keep_together = True
        elif p.style.name == 'Caption':
            p.paragraph_format.keep_together = True
            p.paragraph_format.keep_with_next = False
    # Present the useful top region of the JSON viewport, leaving original media intact.
    for i, p in enumerate(doc.paragraphs):
        if p.text.startswith('E-P7-02') and i > 0:
            for inline in doc.paragraphs[i-1]._p.xpath('.//wp:inline'):
                ns_a = 'http://schemas.openxmlformats.org/drawingml/2006/main'
                ns_pic = 'http://schemas.openxmlformats.org/drawingml/2006/picture'
                fill = inline.find('.//{' + ns_pic + '}blipFill')
                if fill is not None:
                    rect = fill.find('{' + ns_a + '}srcRect')
                    if rect is None:
                        rect = OxmlElement('a:srcRect'); fill.insert(1, rect)
                    rect.set('b', '90000')
                    extent = inline.find(qn('wp:extent'))
                    extent.set('cy', str(round(int(extent.get('cy')) * .10)))
                    ext = inline.find('.//{' + ns_pic + '}spPr/{' + ns_a + '}xfrm/{' + ns_a + '}ext')
                    if ext is not None:
                        ext.set('cy', extent.get('cy'))
            set_text(p, 'E-P7-02 • Trích phần đầu giao diện JSON canary chứa marker thử nghiệm.')
    # Keep the original monochrome Times New Roman 13 design, including generated fields.
    for style in doc.styles:
        if hasattr(style, 'font'):
            style.font.name = 'Times New Roman'
            style.font.size = Pt(13)
            style.font.color.rgb = RGBColor(0, 0, 0)
    for part in [doc.part] + [s.header.part for s in doc.sections] + [s.footer.part for s in doc.sections]:
        for run in part.element.iter(qn('w:r')):
            normalize_run(run)
        for border in list(part.element.iter(qn('w:pBdr'))):
            border.getparent().remove(border)
    for border in list(doc.styles.element.iter(qn('w:pBdr'))):
        border.getparent().remove(border)
    doc.core_properties.author = 'ShopLab'
    doc.core_properties.subject = 'Kết quả kiểm thử bảo mật ứng dụng web ShopLab'
    doc.core_properties.title = 'Báo cáo kiểm thử bảo mật ứng dụng web ShopLab'
    doc.core_properties.comments = ''
    doc.core_properties.keywords = ''
    doc.core_properties.modified = datetime(2026, 10, 6, tzinfo=timezone.utc)
    doc.save(OUTPUT)
    (WORK / 'edit_log.json').write_text(json.dumps(changes, ensure_ascii=False, indent=2), encoding='utf-8')
    print(json.dumps({'output': str(OUTPUT), 'changes': len(changes), 'paragraphs': len(doc.paragraphs), 'tables': len(doc.tables), 'images': len(doc.inline_shapes)}, ensure_ascii=False))

if __name__ == '__main__':
    main()
