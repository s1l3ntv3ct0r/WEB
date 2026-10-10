from pathlib import Path
from docx import Document
from docx.shared import Pt, RGBColor, Inches
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

root=Path(__file__).resolve().parents[1]
source=root/'pentest/reports/Bao_cao_demo_Brute_force_CAPTCHA_ShopLab.docx'
out=source.with_name('Bao_cao_demo_Brute_force_CAPTCHA_ShopLab_TNR13.docx')
d=Document(source)
replacements={
'accounts.txt và passwords.txt trên Desktop; 300 dòng mỗi file tại lần kiểm tra':'Danh sách chuẩn bị gồm 300 cặp email và mật khẩu; lượt demo theo IP bị chặn tại cặp thứ 28.',
'Danh sách 300 dòng là một tập cặp thông tin đăng nhập.':'Danh sách đầu vào có 300 cặp, nhưng đây là số lượng chuẩn bị, không phải số cặp đã được kiểm tra mật khẩu trong lượt demo. Lượt chạy có giới hạn IP chỉ tiến tới cặp thứ 28 thì nhận 429; các cặp 1–27 đã nhận kết quả đăng nhập, còn cặp 28 bị chặn trước bước xác thực. Không thể thử liền 300 cặp bằng một IP trong cùng cửa sổ 15 phút với cấu hình hiện tại.',
'Hai ảnh cho thấy danh sách ban đầu gồm các địa chỉ user001, user002 và các email tiếp theo.':'Hai ảnh minh họa danh sách email đầu vào, gồm user001, user002 và các địa chỉ tiếp theo. Tổng tập dữ liệu chuẩn bị là 300 cặp; phạm vi quan sát của lượt demo bị rate limit chỉ đến cặp thứ 28, không phải toàn bộ 300 cặp.',
'Ảnh bổ sung cho thấy dòng 25 còn trả 200 cho test, các dòng 26 và 27 trả 401, đến dòng 28 trả 429.':'Ảnh bổ sung cho thấy dòng 25 còn trả 200 cho test, các dòng 26 và 27 trả 401, đến dòng 28 trả 429. Vì vậy, có thể mô tả lượt demo là “thử đến cặp thứ 28 thì bị giới hạn theo IP”, nhưng không ghi “đã xác thực đủ 28 cặp”: cặp thứ 28 chưa được kiểm tra mật khẩu. Các dòng sau nếu tiếp tục gửi trong thời gian bị giới hạn chỉ nhận 429, không phải kết quả đúng hoặc sai mật khẩu.',
'Chạy bộ 6 dòng trước khi dùng tập lớn.':'Chạy bộ 6 dòng để kiểm tra cấu hình. Với tập 300 cặp đã chuẩn bị, lượt demo trong ảnh bị giới hạn tại cặp thứ 28; không báo cáo đã thử hoàn tất 300 cặp trên một IP trong cùng cửa sổ 15 phút.',
}
paras=list(d.paragraphs)
for t in d.tables:
    for row in t.rows:
        for c in row.cells:paras.extend(c.paragraphs)
for para in paras:
    text=para.text
    for a,b in replacements.items():text=text.replace(a,b)
    if text!=para.text:para.text=text

# Apply the requested typography to all editable text, including styles and fields.
parts=[d.part]
for sec in d.sections:parts.extend([sec.header.part,sec.footer.part])
for style in d.styles:
    if hasattr(style,'font'):
        style.font.name='Times New Roman';style.font.size=Pt(13);style.font.color.rgb=RGBColor(0,0,0)
        rp=style.element.get_or_add_rPr()
        fonts=rp.find(qn('w:rFonts'))
        if fonts is not None:
            for a in list(fonts.attrib):
                if 'Theme' in a:del fonts.attrib[a]
            for a in ['ascii','hAnsi','eastAsia','cs']:fonts.set(qn('w:'+a),'Times New Roman')
for part in parts:
    for run in part.element.iter(qn('w:r')):
        rp=run.find(qn('w:rPr'))
        if rp is None:rp=OxmlElement('w:rPr');run.insert(0,rp)
        for name,attrs in [('rFonts',{'ascii':'Times New Roman','hAnsi':'Times New Roman','eastAsia':'Times New Roman','cs':'Times New Roman'}),('sz',{'val':'26'}),('szCs',{'val':'26'}),('color',{'val':'000000'})]:
            old=rp.find(qn('w:'+name))
            if old is not None:rp.remove(old)
            el=OxmlElement('w:'+name)
            for k,v in attrs.items():el.set(qn('w:'+k),v)
            rp.append(el)
for name in ['Normal','Title','Subtitle','Heading 1','Heading 2','Caption']:
    f=d.styles[name].paragraph_format;f.line_spacing=1.05;f.space_after=Pt(5)
for para in paras:
    para.paragraph_format.line_spacing=1.05
    para.paragraph_format.space_after=Pt(5)
    para.paragraph_format.space_before=Pt(0)
for shape in d.inline_shapes:
    shape.width=int(shape.width*.84);shape.height=int(shape.height*.84)
for t in d.tables:
    if t.cell(0,0).text in ['Dòng','Bước','Tình huống']:
        delta=Inches(.22 if t.cell(0,0).text in ['Dòng','Bước'] else .68)
        widths=[c.width for c in t.columns]
        widths[0]+=delta;widths[1]-=delta
        for c,w in zip(t.columns,widths):c.width=w
        for row in t.rows:
            for c,w in zip(row.cells,widths):c.width=w
d.save(out)
print(out)
