"""Generate the editable ShopLab assessment form with the bundled Python runtime."""
from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parents[1]
doc = Document()
doc.settings.odd_and_even_pages_header_footer = False
for border in list(doc.styles.element.iter(qn('w:pBdr'))):
    border.getparent().remove(border)
for border in list(doc.element.iter(qn('w:pBdr'))):
    border.getparent().remove(border)
s = doc.sections[0]
s.page_width, s.page_height = Inches(8.5), Inches(11)
s.top_margin = s.bottom_margin = Inches(.7)
s.left_margin = s.right_margin = Inches(.8)
for name in ['Normal','Title','Subtitle','Heading 1','Heading 2','Heading 3']:
    st=doc.styles[name]; st.font.name='Calibri'; st.font.color.rgb=RGBColor(0,0,0)
doc.styles['Normal'].font.size=Pt(11)
doc.styles['Normal'].paragraph_format.space_after=Pt(6)
doc.styles['Normal'].paragraph_format.line_spacing=1.08
doc.styles['Title'].font.size=Pt(26)
doc.styles['Heading 1'].font.size=Pt(18)
doc.styles['Heading 2'].font.size=Pt(12)
header=s.header.paragraphs[0]; header.text='SHOPLAB  |  HỒ SƠ ĐÁNH GIÁ BẢO MẬT'; header.style='Caption'
header.runs[0].font.color.rgb=RGBColor(0,0,0)
footer=s.footer.paragraphs[0]; footer.text='Bản nháp chưa có kết quả  •  Trang '
field=OxmlElement('w:fldSimple'); field.set(qn('w:instr'),'PAGE'); footer._p.append(field)

def p(t='',style=None): return doc.add_paragraph(t,style)
def h(t): doc.add_heading(t,2)
def page(t):
    doc.add_page_break(); doc.add_heading(t,1)
def entry(label,hint='[ĐIỀN NỘI DUNG]'):
    a=p(); a.add_run(label+'  ').bold=True; a.add_run(hint)
def table(headers,rows,widths):
    t=doc.add_table(rows=1, cols=len(headers)); t.autofit=False
    for col,w in zip(t.columns,widths): col.width=Inches(w)
    for c,txt,w in zip(t.rows[0].cells,headers,widths): c.text=txt;c.width=Inches(w)
    repeat=OxmlElement('w:tblHeader');t.rows[0]._tr.get_or_add_trPr().append(repeat)
    for row in rows:
        for c,txt,w in zip(t.add_row().cells,row,widths):c.text=txt;c.width=Inches(w)
    for i,row in enumerate(t.rows):
        no=OxmlElement('w:cantSplit');row._tr.get_or_add_trPr().append(no)
        for c in row.cells:
            pr=c._tc.get_or_add_tcPr(); borders=OxmlElement('w:tcBorders')
            for edge in ['top','left','bottom','right']:
                e=OxmlElement('w:'+edge);e.set(qn('w:val'),'single');e.set(qn('w:sz'),'4');e.set(qn('w:color'),'D9D9D9');borders.append(e)
            pr.append(borders)
            margins=OxmlElement('w:tcMar')
            for edge in ['top','left','bottom','right']:
                e=OxmlElement('w:'+edge);e.set(qn('w:w'),'85');e.set(qn('w:type'),'dxa');margins.append(e)
            pr.append(margins)
            shade=OxmlElement('w:shd');shade.set(qn('w:fill'),'E8EEF4' if i==0 else 'FFFFFF');pr.append(shade)
            for para in c.paragraphs:
                para.paragraph_format.space_after=Pt(2)
                for r in para.runs:r.font.size=Pt(10);r.bold=i==0
    p()

doc.add_heading('Báo cáo kiểm thử bảo mật ShopLab',0)
p('Mẫu điền trong quá trình pentest', 'Subtitle')
p('Tài liệu này dùng để ghi phạm vi, kết quả kiểm thử, phát hiện đã xác nhận và kết quả kiểm tra sau sửa. Các ô [ĐIỀN] là nội dung cần bổ sung. Chưa có kết quả bảo mật nào được xác nhận trong mẫu.')
entry('Mã đợt đánh giá','SL-PT-[YYYYMM]-[NN]')
entry('Người thực hiện','[ĐIỀN tên]')
entry('Người nhận','[ĐIỀN chủ dự án hoặc người đánh giá]')
entry('Thời gian kiểm thử','[ĐIỀN ngày bắt đầu và kết thúc, múi giờ]')
entry('Phiên bản ứng dụng','[ĐIỀN commit SHA và manifest nếu working tree chưa sạch]')
entry('Phiên bản báo cáo','0.1 Draft')
entry('Phân loại','Internal; tạo bản riêng đã che dữ liệu khi đưa vào portfolio')
h('Cách sử dụng')
for x in ['Điền phạm vi và môi trường trước buổi test đầu tiên.', 'Sau mỗi case, điền nhật ký ở phần 8 và cập nhật coverage ở phần 4.', 'Chỉ nhân bản phần 5 và 6 thành finding khi đã có bằng chứng xác nhận.', 'Viết tóm tắt ở phần 1 sau khi chốt kết quả; cập nhật phần 7 khi retest.', 'Trước bàn giao, bỏ hướng dẫn điền, xóa hàng mẫu chưa dùng và đổi footer bản nháp.']:
    p(x,'List Bullet')
h('Danh mục nội dung')
p('1 Tóm tắt và lịch sử • 2 Phạm vi • 3 Môi trường và phương pháp • 4 Coverage và findings • 5 Chi tiết finding • 6 Bằng chứng và khắc phục • 7 Retest • 8 Nhật ký • 9 Bàn giao và nguồn')

page('1 Tóm tắt và lịch sử báo cáo')
h('Tóm tắt cho chủ dự án')
p('[ĐIỀN SAU KHI TEST: Trong thời gian …, tôi đánh giá ShopLab build … với phạm vi …. Tôi đã xác nhận … phát hiện. Tác động chính đã chứng minh là …. Các giới hạn quan trọng gồm …. Ba ưu tiên xử lý là ….]')
p('Nếu chưa xác nhận lỗi, ghi rõ số case đã thực hiện và giới hạn; không kết luận toàn hệ thống an toàn.')
table(['Critical','High','Medium','Low','Thông tin'],[['[Chưa chốt]']*5],[1.38]*5)
entry('Rủi ro đã chứng minh','[ĐIỀN dữ liệu hoặc nghiệp vụ bị ảnh hưởng]')
entry('Điểm kiểm soát hoạt động đúng','[ĐIỀN case và evidence tương ứng]')
entry('Ưu tiên 1','[ĐIỀN hành động và người phụ trách]')
entry('Ưu tiên 2','[ĐIỀN]')
entry('Ưu tiên 3','[ĐIỀN]')
h('Lịch sử tài liệu')
table(['Phiên bản','Ngày','Người sửa','Nội dung thay đổi'],[['0.1','[ĐIỀN]','[ĐIỀN]','Khởi tạo hồ sơ'],['[ĐIỀN]','[ĐIỀN]','[ĐIỀN]','[ĐIỀN]']],[.8,1,1.3,3.8])
entry('Kết luận tại thời điểm báo cáo','[ĐIỀN sau khi kiểm thử và rà soát evidence]')

page('2 Phạm vi và quy tắc thực hiện')
table(['Tài sản','Phạm vi dự kiến','Xác nhận thực tế'],[
['localhost:5173','UI và API qua /api','[ĐIỀN build và ngày]'],
['localhost:4000','API trực tiếp nếu đúng tiến trình ShopLab','[ĐIỀN]'],
['localhost:8080','Nginx khi môi trường Docker đã sẵn sàng','Chưa xác nhận'],
['Fixture 4101 và 4102','URL import trong mạng thử nghiệm được kiểm soát','Chưa xác nhận biên mạng'],
['Inbox localhost:8025','Đọc email reset tài khoản test; không scan inbox','[ĐIỀN]']],[1.4,3.7,1.8])
entry('Người xác nhận phạm vi','[ĐIỀN chủ dự án và ngày xác nhận]')
entry('Chức năng trong phạm vi','Auth, profile, sản phẩm/review, cart/coupon, order, ticket/attachment, admin. [SỬA theo thực tế]')
entry('Ngoài phạm vi','Internet và dịch vụ bên thứ ba, host OS, container escape, DoS, social engineering, dữ liệu thật.')
entry('Giới hạn thao tác','[ĐIỀN tốc độ request, số request đồng thời, giới hạn thử login, thời gian test]')
entry('Điều kiện dừng','[ĐIỀN lỗi dịch vụ, tác động ngoài fixture hoặc request vượt scope]')
entry('Sao lưu và phục hồi','[ĐIỀN vị trí riêng, thời điểm, người quản lý và lần kiểm tra phục hồi; không ghi secret]')
entry('Giới hạn chưa kiểm thử','[ĐIỀN HTTPS, Docker, thiếu quyền, thiếu fixture hoặc thời gian]')
p('Phạm vi dự kiến phải được xác nhận trước khi kiểm thử. Cổng khác trên localhost không tự động là tài sản ShopLab.')

page('3 Môi trường và phương pháp')
entry('Hình thức','Tự đánh giá gray-box với tài khoản được cấp; HTTP trước, đối chiếu source sau. Không phải đánh giá độc lập.')
entry('Stack dự kiến','React, Express, PostgreSQL, Prisma. [ĐIỀN phiên bản thực tế]')
entry('Build và dữ liệu','[ĐIỀN commit/hash manifest, migration, phiên bản fixture]')
entry('Cấu hình đáng chú ý','[ĐIỀN origin, HTTP/HTTPS, cookie mode, proxy; che secret]')
table(['Danh tính','Vai trò','Dữ liệu sở hữu'],[['Anonymous','Chưa đăng nhập','Không có'],['[User A]','Khách hàng','[Order, ticket, item IDs]'],['[User B]','Khách hàng','[Order, ticket, item IDs]'],['[Admin]','Quản trị','[Fixture quản lý]']],[1.5,1.5,3.9])
table(['Công cụ','Phiên bản','Mục đích thực tế'],[['Burp hoặc ZAP','[ĐIỀN]','[ĐIỀN]'],['Browser và DevTools','[ĐIỀN]','[ĐIỀN]'],['Script hoặc test runner','[ĐIỀN]','[ĐIỀN]']],[2,1.2,3.7])
h('Trình tự thực hiện')
p('Chốt scope → lập inventory → phân tích → xác nhận có đối chứng → báo cáo → sửa và retest. Ghi evidence trong từng bước.')
entry('Chuẩn đối chiếu','WSTG v4.2 cho phương pháp; ghi đúng mã và URL từng mục khi ánh xạ. Checklist SL là checklist riêng của ShopLab.')
entry('Đánh giá mức độ','[ĐIỀN quy tắc severity đã dùng; nếu dùng CVSS, ghi phiên bản, vector và lý do metric. Không đánh đồng độ tin cậy với mức độ tác động.]')

page('4 Coverage và phát hiện tổng hợp')
p('Not run: chưa làm. Pass: case cụ thể đạt. Fail: có lỗi đã xác nhận. Blocked: bị chặn, ghi lý do. N/A: không áp dụng, ghi căn cứ. Nghi vấn chưa xác nhận được theo dõi riêng.')
table(['Nhóm kiểm tra','Trạng thái','Case và evidence'],[[x,'Not run','[ĐIỀN]'] for x in ['Scope và mapping','Authentication và session','Authorization ngang và dọc','Cart coupon order và race','Input và browser','Upload download URL import','Cấu hình dependency source']],[2.8,1.1,3])
entry('Operation coverage','[ĐIỀN số method/path/role đã test trên tổng inventory; dẫn file inventory và checklist. Không tính Blocked thành Pass.]')
h('Phát hiện đã xác nhận')
table(['ID','Tên vấn đề cụ thể','Mức độ','Trạng thái'],[['[F-xxx]','[CHỈ THÊM SAU XÁC NHẬN]','[ĐIỀN]','[ĐIỀN]']],[.7,3.5,1.1,1.6])
h('Nghi vấn và giới hạn cần làm tiếp')
table(['Mã theo dõi','Điều chưa rõ','Bước xác minh tiếp'],[['[U-xxx]','[ĐIỀN; không đưa vào số findings]','[ĐIỀN]']],[1,3,2.9])

page('5 Chi tiết finding để nhân bản')
p('Nhân bản phần 5 và 6 cho mỗi finding đã xác nhận. Giữ nguyên ID qua các lần retest.')
entry('Finding ID và tên','F-[NNN] — [ĐIỀN hành vi và tác động cụ thể]')
entry('Trạng thái xác nhận','[HTTP đã tái lập / source-only đã xác nhận; nêu giới hạn]')
entry('Build và ngày','[ĐIỀN]')
entry('Endpoint và method','[ĐIỀN tất cả vị trí bị ảnh hưởng đã xác minh]')
entry('Vai trò và tiền điều kiện','[ĐIỀN token thuộc ai, dữ liệu và cấu hình cần có; không ghi token thật]')
entry('WSTG và CWE','[ĐIỀN đúng phiên bản, mã và URL đã đối chiếu]')
entry('Severity','[ĐIỀN mức độ và căn cứ quyền, điều kiện, phạm vi, tác động]')
entry('CVSS nếu dùng','[ĐIỀN phiên bản, vector, điểm đã tính và giải thích metric; nếu không dùng ghi N/A]')
h('Mô tả và tác động')
entry('Hành vi mong đợi','[ĐIỀN kiểm soát cần có]')
entry('Hành vi quan sát','[ĐIỀN điều thực sự xảy ra]')
entry('Tác động đã chứng minh','[ĐIỀN dữ liệu đọc/sửa, chủ sở hữu và phạm vi; tách giả thuyết mở rộng]')
h('Các bước tái lập')
for label in ['1 Chuẩn bị fixture và xác nhận chủ sở hữu','2 Gửi request hợp lệ làm đối chứng','3 Thay token ID field hoặc trạng thái','4 Ghi response và kiểm tra dữ liệu sau thao tác','5 Lặp lại trên fixture mới để xác nhận']:
    entry(label,'[ĐIỀN thao tác chính xác]')

page('6 Bằng chứng và khắc phục finding')
entry('Finding liên quan','F-[NNN]')
h('Request và response')
p('[DÁN baseline request đã che Authorization/Cookie và phần response liên quan. Ghi status, requestId và thời gian.]')
p('[DÁN request thay đổi; đánh dấu đúng phần đã sửa và response chứng minh vấn đề. Có thể dẫn tệp text thay vì dán toàn bộ.]')
h('Bằng chứng hình ảnh')
p('[CHÈN ảnh đã cắt đúng vùng, che secret. Ghi chú: Hình F-NNN-01 — thao tác, role, điều ảnh chứng minh. Ảnh không thay thế bản request/response.]')
table(['Evidence ID','Tệp và vị trí lưu','Nội dung và thời gian'],[['E-[NNN]-01','[ĐIỀN]','[ĐIỀN]'],['E-[NNN]-02','[ĐIỀN]','[ĐIỀN]']],[1.1,2.8,3])
entry('Dữ liệu trước và sau','[ĐIỀN owner, order/stock/redemption hoặc quyền truy cập liên quan]')
entry('Nguyên nhân','[ĐIỀN route/service/constraint và dòng mã nếu đã xác nhận]')
entry('Biện pháp sửa','[ĐIỀN kiểm soát cụ thể và vị trí áp dụng; không chỉ ghi validate input]')
entry('Người phụ trách và thời hạn','[ĐIỀN]')
entry('Regression test','[ĐIỀN PoC phải bị chặn, biến thể liên quan và happy path phải còn hoạt động]')

page('7 Kế hoạch sửa và retest')
table(['Finding','Phụ trách và mốc','Commit sửa','Retest'],[['[F-xxx]','[ĐIỀN]','[Chưa có]','Not retested']],[1,2.2,1.6,2.1])
h('Biên bản retest để nhân bản')
entry('Finding','[F-xxx]')
entry('Ngày và người kiểm tra','[ĐIỀN]')
entry('Build ban đầu và build sửa','[ĐIỀN cả hai commit hoặc manifest]')
entry('Thay đổi môi trường','[ĐIỀN; nêu ảnh hưởng đến khả năng so sánh]')
table(['Phép thử','Expected','Actual và evidence'],[['PoC ban đầu','[ĐIỀN]','[Chưa thực hiện]'],['Biến thể lân cận','[ĐIỀN]','[Chưa thực hiện]'],['Thao tác hợp lệ','[ĐIỀN]','[Chưa thực hiện]']],[1.5,2.3,3.1])
entry('Kết quả','Not retested [đổi thành Fixed / Partially fixed / Not fixed sau kiểm tra]')
entry('Rủi ro còn lại','[ĐIỀN điều chưa xử lý hoặc chưa thể xác minh]')
entry('Evidence retest','[ĐIỀN tệp mới; không ghi đè bằng chứng ban đầu]')
entry('Quyết định tiếp theo','[ĐIỀN đóng finding hoặc hành động tiếp theo; chấp nhận rủi ro không đồng nghĩa đã sửa]')

page('8 Nhật ký từng case để nhân bản')
p('Phần làm việc: cập nhật ngay khi test. Khi bàn giao có thể chuyển nhật ký đầy đủ thành phụ lục riêng, giữ báo cáo chính gọn.')
entry('Buổi và case','[ĐIỀN buổi 01–10; SL-xx; biến thể cụ thể]')
entry('Ngày giờ build người test','[ĐIỀN]')
entry('Giả thuyết','[ĐIỀN điều cần kiểm tra, không viết như lỗi đã có]')
entry('Role và tài nguyên','[ĐIỀN A/B/admin/anonymous; owner và ID]')
entry('Request chuẩn','[ĐIỀN method/path và evidence baseline]')
entry('Thay đổi trong phép thử','[ĐIỀN một yếu tố hoặc tổ hợp có giải thích]')
entry('Expected','[ĐIỀN status, nội dung và dữ liệu phải giữ nguyên]')
entry('Actual','[ĐIỀN status, nội dung và dữ liệu thực tế trước/sau]')
entry('Đối chứng','[ĐIỀN token còn hợp lệ, đúng owner, không bị rate limit hoặc dữ liệu động gây nhiễu]')
entry('Evidence','[ĐIỀN đường dẫn đã che secret và requestId]')
entry('Kết quả','Not run [cập nhật Pass / Fail / Blocked / N/A; nghi vấn chưa đủ bằng chứng giữ In progress trong notes]')
entry('Finding liên quan','[F-xxx nếu đã xác nhận; nếu không ghi Không có]')
entry('Dọn fixture','[ĐIỀN đúng bản ghi/tệp đã tạo và việc còn cần giữ]')
entry('Việc tiếp theo','[ĐIỀN bước còn thiếu và người thực hiện]')
h('Tổng kết cuối buổi')
entry('Đã làm và chưa làm','[ĐIỀN case IDs, blockers và việc ưu tiên buổi tới]')

page('9 Kiểm tra bàn giao và nguồn')
for x in ['Build, ngày, phạm vi và danh tính test đã được ghi đúng.', 'Mỗi finding có evidence tái lập và tác động cụ thể; không có nghi vấn bị tính như lỗi xác nhận.', 'Các số liệu summary khớp bảng findings; severity và CVSS không mâu thuẫn.', 'Coverage ghi rõ Pass, Fail, Not run, Blocked và N/A với lý do.', 'Token, Cookie, reset link, secret và dữ liệu riêng đã được che.', 'Retest phản ánh đúng việc đã làm; không tự chuyển thành Fixed.', 'Ảnh đọc được, bảng không vỡ, các liên kết evidence hoạt động.', 'Đã bỏ hướng dẫn và placeholder chưa dùng trước khi bàn giao bản cuối.']:
    p('[ ] '+x)
h('Danh mục đính kèm')
table(['Hồ sơ','Vị trí hoặc liên kết','Tình trạng'],[['Inventory method path role','[ĐIỀN]','Chưa đính kèm'],['Checklist đã cập nhật','[ĐIỀN]','Chưa đính kèm'],['Evidence đã che secret','[ĐIỀN]','Chưa đính kèm'],['Manifest và cấu hình công cụ','[ĐIỀN]','Chưa đính kèm'],['Regression và retest','[ĐIỀN]','Chưa đính kèm']],[2.5,2.8,1.6])
h('Nguồn cấu trúc')
p('OWASP WSTG v4.2 Reporting\nhttps://wstg.owasp.org/v4.2/5-Reporting/')
p('Mẫu được tùy biến cho ShopLab từ cấu trúc đã đối chiếu ngày 23/09/2026. Đây không phải biểu mẫu do OWASP phát hành hoặc chứng nhận. Các ID SL và F là mã nội bộ của đợt đánh giá.')
p('Tài liệu đồng hành trong dự án: pentest/MASTER_PLAN.md, STEP_BY_STEP.md, CHECKLIST.md và scope.md.')

out=ROOT/'pentest'/'reports'/'ShopLab_Pentest_Report_Form.docx'
out.parent.mkdir(parents=True,exist_ok=True)
doc.save(out)
print(out)
